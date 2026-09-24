from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from abc import ABC, abstractmethod
from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfReader

from ..models.document import ExtractedDocument


class DocumentExtractor(ABC):
    @abstractmethod
    def extract(self, file_bytes: bytes) -> ExtractedDocument:
        """Return normalized extracted content without leaking parser-specific objects."""


class TextExtractor(DocumentExtractor):
    def extract(self, file_bytes: bytes) -> ExtractedDocument:
        return self.extract_text(file_bytes.decode("utf-8", errors="ignore"))

    def extract_text(self, text: str) -> ExtractedDocument:
        cleaned = text.strip()
        return ExtractedDocument(
            text=cleaned,
            page_count=1,
            content_type="text/plain",
            metadata={"source": "text"},
        )


class PdfExtractor(DocumentExtractor):
    def extract(self, file_bytes: bytes) -> ExtractedDocument:
        return self.extract_file_content(file_bytes)

    def extract_file_content(self, file_bytes: bytes) -> ExtractedDocument:
        try:
            reader = PdfReader(BytesIO(file_bytes))
            pages = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                pages.append(page_text)
            text = "\n\n".join(pages).strip()
            return ExtractedDocument(
                text=text,
                page_count=len(pages),
                content_type="application/pdf",
                metadata={"source": "pdf"},
            )
        except Exception:
            return ExtractedDocument(text="", page_count=0, content_type="application/pdf", metadata={"source": "pdf", "error": "parse_failed"})


class DocxExtractor(DocumentExtractor):
    def extract(self, file_bytes: bytes) -> ExtractedDocument:
        try:
            document = DocxDocument(BytesIO(file_bytes))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
            return ExtractedDocument(
                text=text,
                page_count=1,
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                metadata={"source": "docx"},
            )
        except Exception:
            return ExtractedDocument(text="", page_count=0, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", metadata={"source": "docx", "error": "parse_failed"})

    def extract_text_from_xml(self, xml_string: str) -> str:
        pattern = re.compile(r"<(?:[A-Za-z0-9_]+:)?t(?:\s[^>]*)?>(.*?)</(?:[A-Za-z0-9_]+:)?t>", re.DOTALL)
        matches = pattern.findall(xml_string)
        if matches:
            return "".join(matches)

        try:
            root = ET.fromstring(xml_string)
            texts: list[str] = []
            for node in root.iter():
                tag = node.tag.rsplit("}", 1)[-1] if "}" in node.tag else node.tag
                local_tag = tag.split(":", 1)[-1]
                if local_tag == "t" and node.text:
                    texts.append(node.text)
            return "".join(texts)
        except ET.ParseError:
            return ""


class ExtractionService:
    def __init__(self):
        self.extractors = {
            "txt": TextExtractor(),
            "pdf": PdfExtractor(),
            "docx": DocxExtractor(),
        }

    def extract(self, file_type: str, file_bytes: bytes) -> ExtractedDocument:
        extractor = self.extractors.get(file_type.lower())
        if extractor is None:
            raise ValueError(f"Unsupported file type: {file_type}")
        return extractor.extract(file_bytes)
